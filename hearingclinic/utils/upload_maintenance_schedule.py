"""
CONSOLE SCRIPT: Import Schedules without pandas
Uses openpyxl which is included with Frappe/ERPNext

Copy and paste this entire block into bench console

N.B.: This script directly inserts into the database, bypassing Frappe's ORM and validation.
This was necessary due to issues with the validation which did not allow to load own 
"ID (Schedules)" - which we need to load maintenance visits.

"""
import frappe
from openpyxl import load_workbook
from collections import defaultdict

file_path = '/home/frappe/frappe-bench/backups/OUT_08_schedule_item_load.xlsx'
wb = load_workbook(file_path)
ws = wb.active
headers = [cell.value for cell in ws[1]]
col_map = {header: idx for idx, header in enumerate(headers)}

schedule_groups = defaultdict(list)
item_end_dates = {}
current_parent_id = None

for row in ws.iter_rows(min_row=2, values_only=True):
    parent_id = row[col_map['ID']]
    if parent_id is not None:
        current_parent_id = parent_id
        if 'End Date (Items)' in col_map and row[col_map['End Date (Items)']] is not None:
            item_end_dates[parent_id] = row[col_map['End Date (Items)']]
    
    if current_parent_id and row[col_map['Scheduled Date (Schedules)']] is not None:
        schedule_groups[current_parent_id].append({
            'name': row[col_map['ID (Schedules)']],
            'scheduled_date': row[col_map['Scheduled Date (Schedules)']],
            'item_code': row[col_map['Item Code (Schedules)']],
            'item_name': row[col_map['Item Name (Schedules)']],
            'item_reference': row[col_map['Item Reference (Schedules)']],
            'completion_status': row[col_map.get('Completion Status (Schedules)')],
            'actual_date': row[col_map.get('Actual Date (Schedules)')]
        })

print(f"Loaded {len(schedule_groups)} parent documents")
print(f"Loaded {len(item_end_dates)} end dates")

success = 0
total = len(schedule_groups)

for parent_id, schedules in schedule_groups.items():
    try:
        # Delete existing schedules
        frappe.db.sql("DELETE FROM `tabMaintenance Schedule Detail` WHERE parent = %s", parent_id)
        
        # Insert new schedules
        for idx, s in enumerate(schedules, 1):
            frappe.db.sql("""
                INSERT INTO `tabMaintenance Schedule Detail` 
                (name, parent, parenttype, parentfield, idx, scheduled_date, item_code, item_name, item_reference, completion_status, actual_date, creation, modified, modified_by, owner, docstatus)
                VALUES (%s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, NOW(), NOW(), %s, %s, 0)
            """, (s['name'], parent_id, 'Maintenance Schedule', 'schedules', idx, s['scheduled_date'], s['item_code'], s['item_name'], s['item_reference'], s['completion_status'], s['actual_date'], frappe.session.user, frappe.session.user))
        
        # Update End Date - try end_date column
        if parent_id in item_end_dates:
            frappe.db.sql("UPDATE `tabMaintenance Schedule Item` SET end_date = %s WHERE parent = %s", (item_end_dates[parent_id], parent_id))
        
        success += 1
        if success % 50 == 0:
            frappe.db.commit()
            print(f"  {success}/{total} processed...")
    
    except Exception as e:
        print(f"Error with {parent_id}: {e}")

frappe.db.commit()
print(f"\nCOMPLETE! {success}/{total} documents imported")